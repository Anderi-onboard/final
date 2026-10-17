"""IR -> .docx / .txt / .md / .html"""
from __future__ import annotations

import base64
import html
import io
import re
from typing import List

from .model import Doc, Image, Page, Para, Run, Table

DEFAULT_LATIN = "Times New Roman"
DEFAULT_EAST = "宋体"
_BAD_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


def _xml(s: str) -> str:
    return _BAD_XML.sub("", s)


# ------------------------------------------------------------------------ docx

def _set_fonts(run_or_style, latin, east):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    el = run_or_style.element if hasattr(run_or_style, "element") else run_or_style._element
    rpr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else el.rPr
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), latin)
    rf.set(qn("w:eastAsia"), east)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(attr)) is not None:
            del rf.attrib[qn(attr)]


def write_docx(doc: Doc, path: str, page_breaks: bool = True):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt, RGBColor

    d = Document()
    normal = d.styles["Normal"]
    normal.font.size = Pt(10.5)
    _set_fonts(normal, DEFAULT_LATIN, DEFAULT_EAST)
    for lvl in (1, 2, 3):
        try:
            _set_fonts(d.styles[f"Heading {lvl}"], DEFAULT_LATIN, DEFAULT_EAST)
        except KeyError:
            pass
    d.core_properties.title = doc.title
    d.core_properties.author = ""

    align = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
             "right": WD_ALIGN_PARAGRAPH.RIGHT}

    def setup(sec, pg: Page):
        sec.page_width, sec.page_height = Pt(pg.width), Pt(pg.height)
        l, t, r, b = pg.margins
        sec.left_margin, sec.top_margin = Pt(l), Pt(t)
        sec.right_margin, sec.bottom_margin = Pt(r), Pt(b)

    cur = None
    for n, pg in enumerate(doc.pages):
        new_section = False
        if n == 0:
            setup(d.sections[0], pg)
            cur = pg
        elif page_breaks and (abs(pg.width - cur.width) > 2 or abs(pg.height - cur.height) > 2
                              or any(abs(a - b) > 2 for a, b in zip(pg.margins, cur.margins))):
            setup(d.add_section(WD_SECTION.NEW_PAGE), pg)   # starts its own page
            cur = pg
            new_section = True
        text_w = max(72.0, cur.width - cur.margins[0] - cur.margins[2])
        want_break = page_breaks and n > 0 and not new_section
        blocks = list(pg.blocks)
        if not blocks and pg.note:
            blocks = [Para([Run(f"[{pg.note}]", 9, italic=True, color="888888")])]
        if want_break:
            if not blocks or not isinstance(blocks[0], Para):
                sp = d.add_paragraph()
                sp.paragraph_format.page_break_before = True
                sp.paragraph_format.space_after = Pt(0)
                sp.paragraph_format.line_spacing = Pt(1)
                want_break = False

        for bi, b in enumerate(blocks):
            if isinstance(b, Para):
                p = d.add_paragraph(style=f"Heading {b.heading}" if b.heading else "Normal")
                pf = p.paragraph_format
                p.alignment = align.get(b.align, WD_ALIGN_PARAGRAPH.LEFT)
                if b.left_indent:
                    pf.left_indent = Pt(b.left_indent)
                if b.first_indent:
                    pf.first_line_indent = Pt(b.first_indent)
                if b.right_indent:
                    pf.right_indent = Pt(b.right_indent)
                pf.space_before = Pt(b.space_before)
                pf.space_after = Pt(b.space_after)
                pf.keep_with_next = False
                if b.line_pitch:
                    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                    pf.line_spacing = Pt(b.line_pitch)
                else:
                    pf.line_spacing = 1.0
                if want_break and bi == 0:
                    pf.page_break_before = True
                for r in b.runs:
                    if not r.text:
                        continue
                    run = p.add_run(_xml(r.text))
                    if r.size:
                        run.font.size = Pt(r.size)
                    run.bold = True if r.bold else (False if b.heading else None)
                    if r.italic:
                        run.italic = True
                    if r.color:
                        run.font.color.rgb = RGBColor.from_string(r.color.upper())
                    elif b.heading:
                        run.font.color.rgb = RGBColor(0, 0, 0)
                    _set_fonts(run, r.font or DEFAULT_LATIN, r.font or DEFAULT_EAST)
            elif isinstance(b, Image):
                w = min(b.width, text_w)
                h = b.height * (w / b.width) if b.width else b.height
                p = d.add_paragraph()
                p.alignment = align.get(b.align, WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.space_before = Pt(b.space_before)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if b.indent and b.align == "left":
                    p.paragraph_format.left_indent = Pt(min(b.indent, max(0, text_w - w)))
                if want_break and bi == 0:
                    p.paragraph_format.page_break_before = True
                try:
                    p.add_run().add_picture(io.BytesIO(b.data), width=Pt(w), height=Pt(h))
                except Exception:
                    pass
            elif isinstance(b, Table):
                _docx_table(d, b, text_w, Pt, WD_TABLE_ALIGNMENT)
    d.save(path)


def _docx_table(d, t: Table, text_w, Pt, WD_TABLE_ALIGNMENT):
    ncol = max(len(r) for r in t.rows)
    tbl = d.add_table(rows=len(t.rows), cols=ncol)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    widths = None
    if t.col_widths and len(t.col_widths) == ncol:
        s = sum(t.col_widths)
        widths = [w * min(1.0, text_w / s) for w in t.col_widths] if s else None
    for ri, row in enumerate(t.rows):
        for ci in range(ncol):
            cell = tbl.cell(ri, ci)
            if widths:
                cell.width = Pt(widths[ci])
            txt = _xml(row[ci]) if ci < len(row) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(txt)
            run.font.size = Pt(10)
            _set_fonts(run, DEFAULT_LATIN, DEFAULT_EAST)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


# ------------------------------------------------------------------ txt / md

def _para_md(p: Para) -> str:
    out = []
    for r in p.runs:
        t = r.text
        if not t.strip():
            out.append(t)
            continue
        lead = t[:len(t) - len(t.lstrip())]
        trail = t[len(t.rstrip()):]
        core = t.strip()
        if r.bold and r.italic:
            core = f"***{core}***"
        elif r.bold and not p.heading:
            core = f"**{core}**"
        elif r.italic:
            core = f"*{core}*"
        out.append(lead + core + trail)
    s = "".join(out).strip()
    return ("#" * p.heading + " " + s) if p.heading else s


def write_txt(doc: Doc, path: str, page_breaks: bool = True):
    parts: List[str] = []
    for n, pg in enumerate(doc.pages):
        if n and page_breaks:
            parts.append("\f")
        for b in pg.blocks:
            if isinstance(b, Para):
                parts.append(b.text.strip())
            elif isinstance(b, Table):
                parts.append("\n".join("\t".join(r) for r in b.rows))
    with open(path, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write("\n\n".join(p for p in parts if p).replace("\n\n\f\n\n", "\n\n\f\n"))
        f.write("\n")


def write_md(doc: Doc, path: str, page_breaks: bool = True):
    parts: List[str] = []
    for n, pg in enumerate(doc.pages):
        if n and page_breaks:
            parts.append("---")
        for b in pg.blocks:
            if isinstance(b, Para):
                parts.append(_para_md(b))
            elif isinstance(b, Table):
                rows = [[c.replace("|", "\\|") for c in r] + [""] * (len(b.rows[0]) - len(r)) for r in b.rows]
                lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * len(rows[0])]
                lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                parts.append("\n".join(lines))
            elif isinstance(b, Image):
                parts.append("*[图片]*")
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n\n".join(p for p in parts if p))
        f.write("\n")


# ------------------------------------------------------------------------ html

def _sniff(data: bytes) -> str:
    if data[:8] == b"\x89PNG\r\n\x1a\n":
        return "image/png"
    if data[:3] == b"\xff\xd8\xff":
        return "image/jpeg"
    if data[:3] == b"GIF":
        return "image/gif"
    return "image/png"


def write_html(doc: Doc, path: str, page_breaks: bool = True):
    css = ("body{font:15px/1.7 'Segoe UI','Microsoft YaHei',sans-serif;max-width:46em;margin:2em auto;"
           "padding:0 1em;color:#222}p{margin:0}table{border-collapse:collapse;margin:.8em 0}"
           "td{border:1px solid #999;padding:.25em .6em;vertical-align:top}"
           "hr.page{border:0;border-top:1px dashed #aaa;margin:2em 0}img{max-width:100%}")
    out = ["<!doctype html><html><head><meta charset='utf-8'>",
           f"<title>{html.escape(doc.title)}</title><style>{css}</style></head><body>"]
    for n, pg in enumerate(doc.pages):
        if n and page_breaks:
            out.append("<hr class='page'>")
        for b in pg.blocks:
            if isinstance(b, Para):
                tag = f"h{b.heading}" if b.heading else "p"
                st = [f"text-align:{b.align}"]
                if b.left_indent:
                    st.append(f"margin-left:{b.left_indent:.0f}pt")
                if b.first_indent:
                    st.append(f"text-indent:{b.first_indent:.0f}pt")
                if b.right_indent:
                    st.append(f"margin-right:{b.right_indent:.0f}pt")
                if b.space_before:
                    st.append(f"margin-top:{b.space_before:.0f}pt")
                spans = []
                for r in b.runs:
                    s = html.escape(r.text)
                    css_r = []
                    if r.size and not b.heading:
                        css_r.append(f"font-size:{r.size}pt")
                    if r.color:
                        css_r.append(f"color:#{r.color}")
                    if r.bold and not b.heading:
                        s = f"<b>{s}</b>"
                    if r.italic:
                        s = f"<i>{s}</i>"
                    spans.append(f"<span style='{';'.join(css_r)}'>{s}</span>" if css_r else s)
                out.append(f"<{tag} style='{';'.join(st)}'>{''.join(spans)}</{tag}>")
            elif isinstance(b, Table):
                rows = "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r) + "</tr>" for r in b.rows)
                out.append(f"<table>{rows}</table>")
            elif isinstance(b, Image):
                out.append(f"<p style='text-align:{b.align}'><img src='data:{_sniff(b.data)};base64,"
                           f"{base64.b64encode(b.data).decode()}'></p>")
    out.append("</body></html>")
    with open(path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(out))


WRITERS = {"docx": write_docx, "txt": write_txt, "md": write_md, "html": write_html}
FORMAT_LABELS = {"docx": "Word (.docx)", "txt": "纯文本 (.txt)", "md": "Markdown (.md)", "html": "网页 (.html)"}
