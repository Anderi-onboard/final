"""python -m unittest discover -s tests -v      (from tools/pdf2word)

No Windows needed: the layout is exercised with synthetic OCR lines, the
Windows engine with a stand-in for the WinRT projection, and the full path with
Tesseract if it happens to be installed."""
import asyncio
import os
import shutil
import sys
import tempfile
import threading
import time
import types
import unittest
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
sys.path.insert(0, HERE)

from p2w import layout_ocr, ocr, writers                      # noqa: E402
from p2w.engine import (CANCELLED, DONE, FAILED, SKIPPED, JobQueue, Options,   # noqa: E402
                        convert_file, parse_range)
from p2w.model import Image, Para, Table                      # noqa: E402
from p2w.ocr import OLine, OWord                               # noqa: E402

try:
    import make_fixtures
    FONT = make_fixtures.cjk_font()
except BaseException:                                          # no CJK font on this box
    make_fixtures, FONT = None, None

try:
    import pytesseract  # noqa: F401
    HAVE_TESS = shutil.which("tesseract") is not None
except ImportError:
    HAVE_TESS = False
try:
    import rapidocr  # noqa: F401
    HAVE_RAPID = True
except ImportError:
    try:
        import rapidocr_onnxruntime  # noqa: F401
        HAVE_RAPID = True
    except ImportError:
        HAVE_RAPID = False


def line(text, x, y, h, cw=None):
    """A synthetic OCR line: words laid out left to right, `cw` px per character."""
    cw = cw or h * 0.5
    words, cx = [], x
    for w in text.split(" "):
        wd = len(w) * cw
        words.append(OWord(w, cx, y, cx + wd, y + h))
        cx += wd + cw * 0.6
    return OLine(words)


def page_lines():
    L = []
    L.append(line("Annual Report", 280, 100, 60, 34))                        # title, centred
    y = 220
    for n, t in enumerate(["The company grew steadily during the year and opened",
                           "three new offices in the region while keeping costs",
                           "under control through a careful hiring plan."]):
        L.append(line(t, 148 if n == 0 else 100, y, 24, 12))                  # first-line indent
        y += 38
    y += 30                                                                  # paragraph gap
    for t in ["A second paragraph follows after a visible gap and",
              "wraps onto one more line of ordinary body text."]:
        L.append(line(t, 100, y, 24, 12))
        y += 38
    y += 30
    for r, row in enumerate([["Item", "Q1", "Q2"], ["Revenue", "120", "135"],
                             ["Cost", "80", "82"], ["Profit", "40", "53"]]):
        for c, cell in enumerate(row):
            L.append(line(cell, 100 + c * 300, y, 24, 12))
        y += 40
    return L


class OcrLayout(unittest.TestCase):
    def setUp(self):
        self.page = layout_ocr.build_page(page_lines(), 1000, 1400, 0.5)

    def test_blocks(self):
        kinds = [type(b).__name__ for b in self.page.blocks]
        self.assertEqual(kinds, ["Para", "Para", "Para", "Table"], kinds)

    def test_title_is_a_centred_heading(self):
        t = self.page.blocks[0]
        self.assertEqual(t.text, "Annual Report")
        self.assertEqual(t.heading, 1)
        self.assertEqual(t.align, "center")
        self.assertGreater(t.runs[0].size, 18)

    def test_paragraphs_merge_lines_and_keep_indent(self):
        p = self.page.blocks[1]
        self.assertTrue(p.text.startswith("The company grew") and p.text.endswith("hiring plan."))
        self.assertEqual(p.text.count("\n"), 0)
        self.assertGreater(p.first_indent, 10)
        self.assertEqual(p.heading, 0)
        self.assertAlmostEqual(p.runs[0].size, self.page.blocks[2].runs[0].size)
        self.assertGreater(self.page.blocks[2].space_before, 0)

    def test_table(self):
        t = self.page.blocks[3]
        self.assertEqual(t.rows[0], ["Item", "Q1", "Q2"])
        self.assertEqual(t.rows[3], ["Profit", "40", "53"])
        self.assertEqual(len(t.col_widths), 3)

    def test_two_columns_read_left_then_right(self):
        L = []
        for i in range(8):
            L.append(line(f"left column text number {i} of the page", 100, 200 + i * 38, 24, 12))
            L.append(line(f"right column text number {i} of the page", 560, 200 + i * 38, 24, 12))
        L.append(line("A Full Width Title", 300, 100, 48, 28))
        pg = layout_ocr.build_page(L, 1000, 1400, 0.5)
        texts = [b.text for b in pg.blocks if isinstance(b, Para)]
        self.assertEqual(texts[0], "A Full Width Title")
        self.assertIn("left column text number 0", texts[1])
        self.assertIn("number 7 of the page", texts[1])
        self.assertIn("right column text number 0", texts[2])

    def test_prose_columns_are_not_a_table(self):
        L = []
        for i in range(8):
            L.append(line("this is a long line of ordinary prose in the left column " + str(i), 100, 200 + i * 38, 24, 7.5))
            L.append(line("this is a long line of ordinary prose in the right col " + str(i), 520, 200 + i * 38, 24, 7.5))
        pg = layout_ocr.build_page(L, 1000, 1400, 0.5)
        self.assertFalse(any(isinstance(b, Table) for b in pg.blocks))

    def test_narrow_book_columns_are_prose_not_a_table(self):
        """Two (and four) columns of Chinese sentences, each only ~30% of the sheet wide:
        the geometry looks like a table, the cell text says otherwise."""
        zh = ["检测既省力又快。这不能说我们的传统文化没用了。", "经过几十年研究,我将奇门遁甲的精髓及经验整理",
              "如下,研究是困难的,但本人不想将已经掌握,经过实", "践得出的一些经验带进坟墓,只想留给后人去研究、去"]
        L = []
        for i in range(8):
            L.append(line(zh[i % 4], 100, 200 + i * 40, 24, 12.5))        # 24 chars * 12.5px = 300px wide
            L.append(line(zh[(i + 1) % 4], 600, 200 + i * 40, 24, 12.5))
        pg = layout_ocr.build_page(L, 1100, 1400, 0.5)
        self.assertFalse(any(isinstance(b, Table) for b in pg.blocks), [type(b).__name__ for b in pg.blocks])
        texts = [b.text for b in pg.blocks if isinstance(b, Para)]
        self.assertEqual(len(texts), 2)                                    # left column, then right column
        self.assertTrue(texts[0].startswith("检测既省力又快"))

    def test_short_label_tables_still_are_tables(self):
        L = []
        for r, row in enumerate([["项目", "一月", "二月", "三月"], ["收入", "120", "135", "150"],
                                 ["成本", "80", "82", "90"], ["利润", "40", "53", "60"]]):
            for c, cell in enumerate(row):
                L.append(line(cell, 100 + c * 220, 300 + r * 40, 24, 12))
        pg = layout_ocr.build_page(L, 1100, 1400, 0.5)
        self.assertTrue(any(isinstance(b, Table) for b in pg.blocks))

    def test_cjk_and_latin_spacing(self):
        l = OLine([OWord("Report", 0, 0, 60, 20), OWord("季度", 70, 0, 110, 20),
                   OWord("报告", 110, 0, 150, 20), OWord("2024", 151, 0, 190, 20)])
        self.assertEqual(l.text, "Report 季度报告2024")

    def test_empty_page(self):
        pg = layout_ocr.build_page([], 800, 1000, 0.5)
        self.assertEqual(pg.blocks, [])
        self.assertTrue(pg.note)


class Writers(unittest.TestCase):
    def test_all_formats(self):
        page = layout_ocr.build_page(page_lines(), 1000, 1400, 0.5)
        from p2w.model import Doc
        doc = Doc("t", [page, layout_ocr.build_page(page_lines(), 1000, 1400, 0.5)])
        with tempfile.TemporaryDirectory() as d:
            for fmt, fn in writers.WRITERS.items():
                path = os.path.join(d, "o." + fmt)
                fn(doc, path, True)
                self.assertGreater(os.path.getsize(path), 100, fmt)
            import docx
            dd = docx.Document(os.path.join(d, "o.docx"))
            self.assertEqual(len(dd.tables), 2)
            self.assertTrue(dd.paragraphs[0].style.name.startswith("Heading"))
            md = open(os.path.join(d, "o.md"), encoding="utf-8").read()
            self.assertIn("| Item | Q1 | Q2 |", md)
            self.assertIn("\n---\n", md)
            self.assertIn("# Annual Report", md)

    def test_control_characters_do_not_break_docx(self):
        from p2w.model import Doc, Page, Run
        pg = Page(595, 842, [Para([Run("bad\x00char\x0b and ok")])])
        with tempfile.TemporaryDirectory() as d:
            writers.write_docx(Doc("t", [pg]), os.path.join(d, "x.docx"))


class Range(unittest.TestCase):
    def test_parse(self):
        self.assertEqual(parse_range("", 4), [0, 1, 2, 3])
        self.assertEqual(parse_range("2-3, 1", 5), [1, 2, 0])
        self.assertEqual(parse_range("3-", 5), [2, 3, 4])
        self.assertEqual(parse_range("1-99", 3), [0, 1, 2])
        with self.assertRaises(ValueError):
            parse_range("x", 3)
        with self.assertRaises(ValueError):
            parse_range("5", 3)


@unittest.skipUnless(FONT, "needs a CJK font to build fixtures")
class Native(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        make_fixtures.make(cls.tmp)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def convert(self, name, **kw):
        out = tempfile.mkdtemp(dir=self.tmp)
        o = Options(out_dir=out, formats=("docx", "md"), collision="overwrite", **kw)
        res = convert_file(os.path.join(self.tmp, name), o, lambda: ocr.make_backend(o.backend, o.ocr_lang))
        return res, out

    def test_text_layer_keeps_formatting(self):
        import docx
        res, out = self.convert("native.pdf")
        self.assertEqual(res["ocr_pages"], 0)
        d = docx.Document(res["outputs"][0])
        styles = [p.style.name for p in d.paragraphs]
        self.assertEqual(styles[0], "Heading 1")
        self.assertIn("Heading 2", styles)
        self.assertIn("季度报告", d.paragraphs[0].text)
        self.assertEqual(len(d.tables), 1)
        self.assertEqual([c.text for c in d.tables[0].rows[1].cells], ["Revenue 收入", "120", "135", "150"])
        self.assertEqual(len(d.inline_shapes), 1)
        self.assertEqual(len(d.sections), 1)
        self.assertTrue(any(p.paragraph_format.page_break_before for p in d.paragraphs))
        sizes = {r.font.size.pt for p in d.paragraphs for r in p.runs if r.font.size}
        self.assertTrue({24.0, 15.0, 11.0} <= sizes, sizes)

    def test_two_columns_stay_in_order_and_narrow(self):
        import docx
        res, _ = self.convert("native.pdf")
        d = docx.Document(res["outputs"][0])
        body = [p for p in d.paragraphs if p.text.startswith(("The committee", "The board"))]
        self.assertEqual([p.text[:9] for p in body][-2:], ["The commi", "The board"])
        left, right = body[-2].paragraph_format, body[-1].paragraph_format
        self.assertIsNotNone(left.right_indent)          # left column must not run under the right one
        self.assertGreater(right.left_indent, left.left_indent or 0)

    def test_no_images_option(self):
        import docx
        res, _ = self.convert("native.pdf", keep_images=False)
        self.assertEqual(len(docx.Document(res["outputs"][0]).inline_shapes), 0)

    def test_page_range_and_no_breaks(self):
        import docx
        res, _ = self.convert("native.pdf", page_range="2", page_breaks=False)
        self.assertEqual(res["pages"], 1)

    def test_native_mode_on_scan_is_empty_not_a_crash(self):
        res, _ = self.convert("scan.pdf", mode="native")
        self.assertEqual(res["ocr_pages"], 0)

    def heads(self, res):
        import docx
        d = docx.Document(res["outputs"][0])
        return [p.text for p in d.paragraphs if p.style.name.startswith("Heading")], d

    def test_split_spread_left_first(self):
        res, _ = self.convert("spread.pdf", split="lr")
        self.assertEqual(res["pages"], 2)
        heads, d = self.heads(res)
        self.assertEqual(heads, ["First Page 第一页", "Second Page 第二页"])
        self.assertEqual(len(d.sections), 1)
        self.assertTrue(any(p.paragraph_format.page_break_before for p in d.paragraphs))
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertLess(text.index("first page of the book"), text.index("second page of the book"))
        self.assertNotIn("second page", text.split("Second Page")[0].lower())   # nothing leaked across the fold

    def test_split_spread_right_first(self):
        heads, _ = self.heads(self.convert("spread.pdf", split="rl")[0])
        self.assertEqual(heads, ["Second Page 第二页", "First Page 第一页"])

    def test_split_auto_only_touches_landscape_sheets(self):
        self.assertEqual(self.convert("spread.pdf", split="auto")[0]["pages"], 2)
        self.assertEqual(self.convert("native.pdf", split="auto")[0]["pages"], 2)    # portrait: untouched
        self.assertEqual(self.convert("spread.pdf", split="none")[0]["pages"], 1)

    def test_split_pages_are_half_width(self):
        import docx
        res, _ = self.convert("spread.pdf", split="lr")
        sec = docx.Document(res["outputs"][0]).sections[0]
        self.assertAlmostEqual(sec.page_width.pt, 421, delta=15)

    @unittest.skipUnless(HAVE_RAPID, "rapidocr not installed")
    def test_split_scanned_spread(self):
        import docx
        res, _ = self.convert("spread_scan.pdf", split="lr", backend="rapidocr")
        self.assertEqual((res["pages"], res["ocr_pages"]), (2, 2))
        d = docx.Document(res["outputs"][0])
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertLess(text.index("First Page"), text.index("Second Page"))
        self.assertLess(text.index("第一页的正文"), text.index("第二页的正文"))

    def test_photo_pages_are_cut_per_photo_then_per_half(self):
        """A portrait page holding two photos of open books: 4 book pages, in reading order,
        cut at the fold (which is not at the middle of the photo), white margins ignored."""
        seen = []

        class Fake:
            name, language = "fake", "x"
            def recognize(self, img):
                seen.append(img.size)
                return []
        o = Options(out_dir=tempfile.mkdtemp(dir=self.tmp), formats=("txt",), collision="overwrite", split="lr")
        res = convert_file(os.path.join(self.tmp, "book.pdf"), o, lambda: Fake())
        self.assertEqual(res["pages"], 4)
        self.assertEqual(len(seen), 4)
        w, h = seen[0][0] + seen[1][0], seen[0][1]
        self.assertEqual(seen[0][1], seen[1][1])
        self.assertAlmostEqual(seen[0][0] / w, 550 / 1200, delta=0.04)      # cut at the fold, not at 50%
        self.assertAlmostEqual(w / h, 1.5, delta=0.1)                       # the photo, not the white page
        for a, b in ((seen[2], seen[0]), (seen[3], seen[1])):
            self.assertLessEqual(max(abs(a[0] - b[0]), abs(a[1] - b[1])), 3)      # same photo shape, rounding aside

    def test_photo_pages_auto_mode_works_on_a_portrait_page(self):
        n = []

        class Fake:
            name, language = "fake", "x"
            def recognize(self, img):
                n.append(1)
                return []
        o = Options(out_dir=tempfile.mkdtemp(dir=self.tmp), formats=("txt",), collision="overwrite", split="auto")
        self.assertEqual(convert_file(os.path.join(self.tmp, "book.pdf"), o, lambda: Fake())["pages"], 4)

    def test_without_split_a_photo_page_stays_one_page(self):
        class Fake:
            name, language = "fake", "x"
            def recognize(self, img):
                return []
        o = Options(out_dir=tempfile.mkdtemp(dir=self.tmp), formats=("txt",), collision="overwrite", split="none")
        self.assertEqual(convert_file(os.path.join(self.tmp, "book.pdf"), o, lambda: Fake())["pages"], 1)

    @unittest.skipUnless(HAVE_RAPID, "rapidocr not installed")
    def test_photo_pages_read_in_order_with_rapidocr(self):
        import docx
        res, _ = self.convert("book.pdf", split="lr", backend="rapidocr")
        self.assertEqual(res["pages"], 4)
        d = docx.Document(res["outputs"][0])
        text = "\n".join(p.text for p in d.paragraphs)
        for a, b in (("甲页第1行", "乙页第1行"), ("乙页第1行", "丙页第1行"), ("丙页第1行", "丁页第1行")):
            self.assertLess(text.index(a), text.index(b), (a, b))
        self.assertGreater(d.sections[0].page_width.pt, 230)                # not a 3-inch page with 5pt type
        sizes = [r.font.size.pt for p in d.paragraphs for r in p.runs if r.font.size]
        self.assertTrue(8 <= sorted(sizes)[len(sizes) // 2] <= 14, sorted(sizes)[len(sizes) // 2])

    def test_spine_is_found_off_centre(self):
        from PIL import Image as PI, ImageDraw
        im = PI.new("L", (1000, 600), 255)
        d = ImageDraw.Draw(im)
        for x0, x1 in ((40, 440), (560, 960)):          # blank strip 440..560 -> fold at 500; text dense elsewhere
            for y in range(40, 560, 12):
                d.rectangle([x0, y, x1, y + 6], fill=0)
        self.assertTrue(440 < ocr.find_spine_px(im) < 560)          # anywhere in the blank strip is a clean cut
        off = PI.new("L", (1000, 600), 255)
        d = ImageDraw.Draw(off)
        for x0, x1 in ((40, 410), (530, 960)):          # gap 410..530 -> fold at 470
            for y in range(40, 560, 12):
                d.rectangle([x0, y, x1, y + 6], fill=0)
        self.assertTrue(410 < ocr.find_spine_px(off) < 530)

    @unittest.skipUnless(HAVE_RAPID, "rapidocr not installed")
    def test_scan_with_rapidocr_gets_text_and_table(self):
        import docx
        res, _ = self.convert("scan.pdf", backend="rapidocr")
        d = docx.Document(res["outputs"][0])
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Quarterly Report", text)
        self.assertIn("委员会详细审阅了方案", text)
        self.assertGreaterEqual(len(d.tables), 1)
        self.assertEqual(d.tables[0].rows[1].cells[0].text, "Revenue 收入")

    @unittest.skipUnless(HAVE_TESS, "tesseract not installed")
    def test_scan_goes_through_ocr(self):
        import docx
        res, _ = self.convert("scan.pdf", backend="tesseract")
        self.assertEqual(res["ocr_pages"], 2)
        d = docx.Document(res["outputs"][0])
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("Quarterly Report", text)
        self.assertIn("committee reviewed the proposal", text)

    @unittest.skipUnless(HAVE_TESS, "tesseract not installed")
    def test_images_and_tiff(self):
        for name, pages in (("photo.jpg", 1), ("multi.tiff", 2)):
            res, _ = self.convert(name, backend="tesseract")
            self.assertEqual(res["pages"], pages, name)
            self.assertEqual(res["ocr_pages"], pages, name)


@unittest.skipUnless(FONT, "needs a CJK font to build fixtures")
class Queue(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp()
        make_fixtures.make(cls.tmp)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def make(self, **kw):
        out = tempfile.mkdtemp(dir=self.tmp)
        kw.setdefault("mode", "native")
        self.opts = Options(out_dir=out, **kw)
        self.idle = threading.Event()
        self.events = []
        q = JobQueue(lambda: self.opts, lambda k, o: (self.idle.set() if k == "idle" else self.events.append((k, o))))
        return q, out

    def wait(self, cond, t=30):
        end = time.time() + t
        while time.time() < end and not cond():
            time.sleep(0.02)
        self.assertTrue(cond(), "timed out")

    def test_batch_runs_in_order_and_goes_idle(self):
        q, out = self.make(workers=2)
        src = os.path.join(self.tmp, "native.pdf")
        copies = []
        for i in range(4):
            c = os.path.join(self.tmp, f"c{i}.pdf"); shutil.copy(src, c); copies.append(c)
        q.add(copies + ["/nonexistent.txt"])           # unsupported types are ignored
        self.assertEqual(len(q.jobs), 4)
        q.start()
        self.assertTrue(self.idle.wait(60))
        self.assertTrue(all(j.status == DONE for j in q.jobs), [(j.status, j.message) for j in q.jobs])
        self.assertEqual(sorted(os.listdir(out)), [f"c{i}.docx" for i in range(4)])

    def test_collision_policies(self):
        q, out = self.make()
        src = os.path.join(self.tmp, "native.pdf")
        for policy, expect in (("rename", ["native (1).docx", "native.docx"]),
                               ("overwrite", ["native.docx"]), ("skip", ["native.docx"])):
            q, out = self.make(collision=policy)
            for _ in range(2):
                self.idle.clear()
                q.add([src]); q.start()
                self.assertTrue(self.idle.wait(30))
            self.assertEqual(sorted(os.listdir(out)), expect, policy)
            if policy == "skip":
                self.assertEqual(q.jobs[-1].status, SKIPPED)

    def test_pause_stop_retry(self):
        import pymupdf
        big = os.path.join(self.tmp, "big.pdf")
        doc = pymupdf.open()
        for i in range(300):
            doc.new_page().insert_text((72, 100), f"page {i} " + "word " * 40)
        doc.save(big)
        q, out = self.make()
        q.add([big])
        q.start(); q.pause()
        time.sleep(0.6)
        held = q.jobs[0].done
        time.sleep(0.6)
        self.assertEqual(q.jobs[0].done, held)          # held at a page boundary
        self.assertLess(held, 300)
        q.stop()
        self.wait(lambda: q.jobs[0].status == CANCELLED)
        self.assertEqual(os.listdir(out), [])
        q.retry()
        self.idle.clear()
        q.start()
        self.assertTrue(self.idle.wait(30))
        self.assertEqual(q.jobs[0].status, DONE)

    def test_bad_files_fail_cleanly_and_do_not_block_the_rest(self):
        q, out = self.make()
        bad = os.path.join(self.tmp, "broken.pdf"); open(bad, "wb").write(b"%PDF-1.4 nope")
        q.add([bad, os.path.join(self.tmp, "native.pdf")])
        q.start()
        self.assertTrue(self.idle.wait(30))
        self.assertEqual([j.status for j in q.jobs], [FAILED, DONE])
        self.assertTrue(q.jobs[0].message)

    def test_ocr_unavailable_fails_only_that_job(self):
        q, out = self.make(mode="auto", backend="windows")        # not Windows here
        q.add([os.path.join(self.tmp, "scan.pdf"), os.path.join(self.tmp, "native.pdf")])
        q.start()
        self.assertTrue(self.idle.wait(30))
        self.assertEqual([j.status for j in q.jobs], [FAILED, DONE])
        self.assertIn("Windows", q.jobs[0].message)


class RapidOcrBackend(unittest.TestCase):
    def test_boxes_become_lines_and_low_confidence_is_dropped(self):
        import numpy as np

        class Out:
            boxes = [[(10, 20), (210, 22), (209, 50), (11, 48)], [(5, 5), (9, 5), (9, 9), (5, 9)]]
            txts = ["Quarterly Report 季度报告", "O"]
            scores = [0.98, 0.3]

        class Engine:
            def __call__(self, arr):
                assert arr.ndim == 3 and arr.shape[2] == 3
                return Out()

        be = ocr.RapidOcr.__new__(ocr.RapidOcr)
        be._new, be._engine, be.language = True, Engine(), "ch+en"
        from PIL import Image as PI
        lines = be.recognize(PI.new("RGB", (300, 100), "white"))
        self.assertEqual([l.text for l in lines], ["Quarterly Report 季度报告"])
        self.assertEqual((lines[0].x0, lines[0].y0, lines[0].x1, lines[0].y1), (10, 20, 210, 50))

    def test_nothing_found(self):
        class Out:
            boxes = txts = scores = None
        be = ocr.RapidOcr.__new__(ocr.RapidOcr)
        be._new, be._engine = True, lambda arr: Out()
        from PIL import Image as PI
        self.assertEqual(be.recognize(PI.new("RGB", (50, 50))), [])

    def test_auto_prefers_rapidocr_when_installed(self):
        try:
            import rapidocr  # noqa: F401
        except ImportError:
            try:
                import rapidocr_onnxruntime  # noqa: F401
            except ImportError:
                self.skipTest("rapidocr not installed")
        self.assertEqual(ocr.make_backend("auto", "auto").name, "rapidocr")


class AppleBackend(unittest.TestCase):
    """Apple Vision needs a Mac. This runs the same calls against a stand-in with the same names
    and shapes (PyObjC style), so the call sequence and the bottom-left -> top-left box
    conversion are covered; whether the real engine reads well is for a Mac to say."""

    def fake_modules(self):
        calls = self.calls = {}
        class Box:
            def __init__(s, x, y, w, h):
                s.origin = types.SimpleNamespace(x=x, y=y); s.size = types.SimpleNamespace(width=w, height=h)
        class Cand:
            def __init__(s, t, c): s._t, s._c = t, c
            def string(s): return s._t
            def confidence(s): return s._c
        class Obs:
            def __init__(s, t, c, box): s.c, s.box = Cand(t, c), box
            def topCandidates_(s, n): return [s.c]
            def boundingBox(s): return s.box
        class Req:
            def init(s): return s
            @classmethod
            def alloc(cls): return cls()
            def setRecognitionLevel_(s, v): calls["level"] = v
            def setUsesLanguageCorrection_(s, v): calls["corr"] = v
            def setRevision_(s, v): calls["rev"] = v
            def setRecognitionLanguages_(s, v): calls["langs"] = v
            def supportedRecognitionLanguagesAndReturnError_(s, e): return (["zh-Hans", "en-US"], None)
            def results(s):
                return [Obs("季度报告 Report", 0.98, Box(0.1, 0.8, 0.5, 0.05)), Obs("噪点", 0.1, Box(0, 0, .1, .1)),
                        Obs("", 0.9, Box(0, 0, .1, .1))]
        class Handler:
            @classmethod
            def alloc(cls): return cls()
            def initWithData_options_(s, d, o): calls["bytes"] = bytes(d); return s
            def performRequests_error_(s, reqs, err): return (True, None)
        class NSData:
            @staticmethod
            def dataWithBytes_length_(b, n): return b
        V = types.SimpleNamespace(VNRecognizeTextRequest=Req, VNImageRequestHandler=Handler,
                                  VNRequestTextRecognitionLevelAccurate=0, VNRecognizeTextRequestRevision3=3)
        return {"Vision": V, "Foundation": types.SimpleNamespace(NSData=NSData)}

    def test_flow_and_coordinates(self):
        from PIL import Image as PI
        with mock.patch.object(ocr.sys, "platform", "darwin"), mock.patch.dict(sys.modules, self.fake_modules()):
            be = ocr.AppleVision("auto")
            self.assertEqual(be.language, "zh-Hans,zh-Hant,en-US")
            lines = be.recognize(PI.new("RGB", (1000, 2000), "white"))
            self.assertEqual([l.text for l in lines], ["季度报告 Report"])          # low confidence + empty dropped
            l = lines[0]
            self.assertEqual((round(l.x0), round(l.x1)), (100, 600))
            self.assertEqual((round(l.y0), round(l.y1)), (300, 400))                 # origin was bottom-left: y flipped
            self.assertEqual(self.calls["level"], 0)
            self.assertEqual(self.calls["rev"], 3)
            self.assertEqual(self.calls["langs"], ["zh-Hans", "zh-Hant", "en-US"])
            self.assertEqual(self.calls["bytes"][:4], b"\x89PNG")
            self.assertEqual([t for t, _ in be.languages()], ["zh-Hans", "en-US"])

    def test_custom_language_and_not_a_mac(self):
        with mock.patch.object(ocr.sys, "platform", "darwin"), mock.patch.dict(sys.modules, self.fake_modules()):
            self.assertEqual(ocr.AppleVision("ja-JP,en-US").langs, ["ja-JP", "en-US"])
        if sys.platform != "darwin":
            with self.assertRaises(ocr.OcrUnavailable):
                ocr.make_backend("apple")

    def test_auto_prefers_apple_on_a_mac(self):
        with mock.patch.object(ocr.sys, "platform", "darwin"), mock.patch.dict(sys.modules, self.fake_modules()):
            self.assertEqual(ocr.make_backend("auto").name, "apple")


class WindowsEngineWithStand_in(unittest.TestCase):
    """Runs WindowsOcr against a fake of the WinRT projection with the same
    names and shapes, so the call sequence and coordinate handling are covered
    even though the real engine only exists on Windows."""

    def setUp(self):
        seen = self.seen = {}

        class Rect:
            def __init__(s, x, y, w, h): s.x, s.y, s.width, s.height = x, y, w, h

        class Word:
            def __init__(s, t, r): s.text, s.bounding_rect = t, r

        class Line:
            def __init__(s, ws): s.words = ws

        class Result:
            lines = [Line([Word("你好", Rect(10, 20, 40, 30)), Word("world", Rect(60, 20, 80, 30))])]

        class Engine:
            recognizer_language = types.SimpleNamespace(language_tag="zh-Hans-CN")
            async def recognize_async(s, bmp):
                seen["bitmap"] = bmp
                return Result()

        class Lang:
            def __init__(s, tag): s.language_tag, s.display_name = tag, tag + " name"

        class OcrEngine:
            available_recognizer_languages = [Lang("en-US"), Lang("zh-Hans-CN")]
            max_image_dimension = 1000
            @staticmethod
            def try_create_from_language(l): seen["lang"] = l.language_tag; return Engine()
            @staticmethod
            def try_create_from_user_profile_languages(): return Engine()

        class Decoder:
            async def get_software_bitmap_async(s): return "BITMAP"
        class BitmapDecoder:
            @staticmethod
            async def create_async(stream): seen["png"] = bytes(stream.data); return Decoder()

        class Stream:
            def __init__(s): s.data = b""
            def seek(s, n): pass
        class Writer:
            def __init__(s, st): s.st = st
            def write_bytes(s, b): s.st.data = bytes(b)
            async def store_async(s): pass
            async def flush_async(s): pass
            def detach_stream(s): pass

        self.w = dict(OcrEngine=OcrEngine, Language=Lang, BitmapDecoder=BitmapDecoder, Stream=Stream, Writer=Writer)

    def test_flow(self):
        from PIL import Image as PI
        with mock.patch.object(ocr, "_WIN", self.w):
            be = ocr.WindowsOcr("auto")
            self.assertEqual(be.language, "zh-Hans-CN")
            self.assertEqual(self.seen["lang"], "zh-Hans-CN")        # Chinese first for mixed text
            lines = be.recognize(PI.new("RGB", (2000, 1000), "white"))     # > max dimension: scaled down
            self.assertEqual(self.seen["png"][:4], b"\x89PNG")
            im = PI.open(__import__("io").BytesIO(self.seen["png"]))
            self.assertLessEqual(max(im.size), 1000)
            # coordinates come back in the pixels of the image we were given (x2)
            self.assertAlmostEqual(lines[0].words[0].x0, 20)
            self.assertAlmostEqual(lines[0].words[1].x1, 280)
            self.assertEqual(lines[0].text, "你好 world")
            self.assertEqual([t for t, _ in be.languages()], ["en-US", "zh-Hans-CN"])

    def test_missing_language_pack_is_explained(self):
        self.w["OcrEngine"].try_create_from_language = staticmethod(lambda l: None)
        with mock.patch.object(ocr, "_WIN", self.w):
            with self.assertRaises(ocr.OcrUnavailable) as cm:
                ocr.WindowsOcr("ja")
            self.assertIn("语言包", str(cm.exception))

    def test_not_windows(self):
        if sys.platform != "win32":
            with self.assertRaises(ocr.OcrUnavailable):
                ocr.make_backend("windows")


if __name__ == "__main__":
    unittest.main()
